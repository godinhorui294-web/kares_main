# -*- coding: utf-8 -*-
"""Verify the generated Word document: structure, numbering, and text sanity."""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm

ROOT = Path(r"C:\kares_main\instances\firefighting_instance")
DOCX = ROOT / "无人机森林灭火任务规划算例说明.docx"
BUILD = ROOT / "_build"

doc = Document(str(DOCX))
paras = [p.text for p in doc.paragraphs]
text = "\n".join(paras)

figs = re.findall(r"图\s*(\d+)　", text)
tabs = re.findall(r"表\s*([A-Z]?\d+)　", text)

print("paragraphs:", len(paras), "tables:", len(doc.tables))
print("inline shapes (pictures):", len(doc.inline_shapes))
print("figure captions:", figs)
print("table captions:", tabs)

dup = [x for x in set(figs) if figs.count(x) > 1]
print("duplicate figure numbers:", dup)
dup_t = [x for x in set(tabs) if tabs.count(x) > 1]
print("duplicate table numbers:", dup_t)
missing = [n for n in range(1, len(figs) + 1) if str(n) not in figs]
print("missing figure numbers:", missing)
seq_t = [str(i) for i in range(1, 20)]
print("missing table numbers 1-19:", [n for n in seq_t if n not in tabs])

bad = [p for p in paras if any(s in p for s in ("None", "nan", "False", "{", "}", "%d", "%s", "if False"))]
print("suspicious paragraphs:", len(bad))
for b in bad[:10]:
    print("   !", b[:120])

# table cell sanity
for i, t in enumerate(doc.tables, 1):
    cells = [c.text for r in t.rows for c in r.cells]
    empties = sum(1 for c in cells if not c.strip())
    width = sum((c.width.cm if c.width else 0) for c in t.rows[0].cells)
    print(f"table {i:2d}: rows={len(t.rows):3d} cols={len(t.columns):2d} empty_cells={empties:3d} "
          f"width={width:.1f}cm")

(BUILD / "docx_text.txt").write_text(text, encoding="utf-8")
print("text dumped to", BUILD / "docx_text.txt")
